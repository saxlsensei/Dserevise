"""從《廉頗藺相如傳》文言字詞總整理 (.docx) 重新生成 js/data.js

用法：python3 tools/build_bank.py "文言字詞總整理.docx"
文件需包含兩個表格：第一個為實詞（編號／例句及字詞／釋義），第二個為虛詞（編號／例句／虛詞／詞性及用法）。
需要：pip install python-docx
"""
import json, random, re, sys, os
import docx

SRC = sys.argv[1] if len(sys.argv) > 1 else "廉頗藺相如列傳_實詞與虛詞總整理.docx"
OUT = os.path.join(os.path.dirname(__file__), "..", "js", "data.js")
rnd = random.Random(20261026)  # 固定種子，確保干擾項穩定，錯題紀錄才能對應

d = docx.Document(SRC)
real = [[c.text.strip() for c in r.cells] for r in d.tables[0].rows[1:]]
func = [[c.text.strip() for c in r.cells] for r in d.tables[1].rows[1:]]

Q, seen = [], {}
pool_real = []
for no, sent, gloss in real:
    k, v = gloss.replace(":", "：").split("：", 1)
    pool_real.append((k, v))

def locate(sentence, k):
    """找出關鍵字在句中第幾次出現（同一句多次出現時按表中次序對應）。"""
    norm = sentence.rstrip("。")
    occ = [m.start() for m in re.finditer(re.escape(k), norm)]
    if not occ:
        return -1
    n = seen.get((norm, k), 0)
    seen[(norm, k)] = n + 1
    return occ[min(n, len(occ) - 1)]

for no, sent, gloss in real:
    k, v = gloss.replace(":", "：").split("：", 1)
    if k == "彊" and "強" in sent:   # 原文字形不一致（彊／強）
        k = "強"
    i = locate(sent, k)
    if i < 0:
        continue
    cands = sorted({vv for kk, vv in pool_real if kk != k and vv != v and kk not in v and k not in vv})
    rnd.shuffle(cands)
    opts = [v] + cands[:3]
    rnd.shuffle(opts)
    Q.append({"id": f"S{len([x for x in Q if x['type']=='real'])+1:03d}", "type": "real", "sentence": sent,
              "hl": [i, len(k)], "word": k, "answer": opts.index(v), "options": opts,
              "explain": f"「{k}」在此句中解作「{v}」。"})

pool_f = sorted({(w, g) for _, _, w, g in func})
for no, sent, w, g in func:
    i = locate(sent, w)
    cands = sorted({gg for ww, gg in pool_f if gg != g and ww != w})
    rnd.shuffle(cands)
    opts = [g] + cands[:3]
    rnd.shuffle(opts)
    Q.append({"id": f"F{len([x for x in Q if x['type']=='func'])+1:03d}", "type": "func", "sentence": sent,
              "hl": [i, len(w)], "word": w, "answer": opts.index(g), "options": opts,
              "explain": f"「{w}」在此句中屬{g}。"})

assert all(len(set(q["options"])) == 4 for q in Q), "干擾項重複"
with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("// 題庫：由文言字詞總整理自動生成（tools/build_bank.py）\nwindow.QBANK = ")
    json.dump(Q, fh, ensure_ascii=False)
    fh.write(";\n")
print(f"已生成 {len(Q)} 題 -> {os.path.normpath(OUT)}")
